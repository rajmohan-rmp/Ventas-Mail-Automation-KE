#!/usr/bin/env python3
"""
ventas_quantity_generator.py — a second, log-scale QUANTITY dashboard (1 slide).

Charts, on a logarithmic Y axis, the daily quantity for:
  Primary Sales  -> TOTAL_QUANTITY
  POS Sales      -> TOTAL_QUANTITY
  Stock Transfer -> REQUESTED_QUANTITY
summed per date across order types, over the same 7-day window as the main deck.

Usage:
  python ventas_quantity_generator.py --csv-file a.csv b.csv ... --date YYYY-MM-DD --output qty.pptx [--png qty.png]
"""
import argparse, sys
from pathlib import Path
from datetime import datetime, timedelta

import ventas_slide_generator as G   # reuse CSV parsing / helpers (same folder)
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

BG="#06224A"; CHART_BG2="#0E3A6E"; HDR="#0E518C"

def _to_float(v):
    try: return float(str(v).strip().replace(",",""))
    except: return 0.0

def _compact(v):
    """Short label so wide numbers don't overlap on narrow log panels: 1.2K / 3.4M / 1.3B."""
    a=abs(v)
    if a>=1e9: return f"{v/1e9:.1f}B"
    if a>=1e6: return f"{v/1e6:.1f}M"
    if a>=1e3: return f"{v/1e3:.1f}K"
    return f"{v:,.0f}" if abs(v-round(v))<1e-9 else f"{v:,.2f}"

def _cfg():
    try:
        import yaml
        return yaml.safe_load((Path(__file__).parent/"config.yaml").read_text(encoding="utf-8")) or {}
    except Exception:
        return {}

def _window(report_date):
    # Mirror ventas_slide_generator.py EXACTLY so both decks in the same email
    # always cover the identical dates: last 7 days, excluding the report day
    # (report_date - 1 back through report_date - 7). The main generator hardcodes
    # this and does not read config, so neither do we.
    end = datetime.strptime(report_date,"%Y-%m-%d") - timedelta(days=1)
    return [(end-timedelta(days=6-i)).strftime("%Y-%m-%d") for i in range(7)]

# metric spec: classify label -> (display title, quantity column, bar colour)
SPEC = {
    "Primary Sales Summary": ("Primary Sales — Total Quantity", "TOTAL_QUANTITY", "#2ECC71"),
    "POS Sales Summary":     ("POS Sales — Total Quantity",     "TOTAL_QUANTITY", "#1172B8"),
    "Stock Transfer Summary":("Stock Transfer — Requested Qty", "REQUESTED_QUANTITY", "#F1C40F"),
}
ORDER = ["Primary Sales Summary","POS Sales Summary","Stock Transfer Summary"]

def _collect(csv_paths):
    """Return {classify_label: {date: quantity}} using the configured column."""
    out={k:{} for k in SPEC}
    for p in csv_paths:
        key=G._classify(p)
        if key not in SPEC: continue
        col=SPEC[key][1]
        rows=G._read_csv_rows(p)
        if not rows: continue
        have = col in rows[0]
        for r in rows:
            d=G._date_key(r)
            if not d: continue
            out[key][d]=out[key].get(d,0.0)+ (_to_float(r.get(col)) if have else 0.0)
    return out

def render_png(data, report_date, out_png, dpi=150):
    dates=_window(report_date)
    short=lambda d: datetime.strptime(d,"%Y-%m-%d").strftime("%d-%b")
    customer=(_cfg().get("branding",{}) or {}).get("customer","")
    x=np.arange(len(dates))
    fig=plt.figure(figsize=(13.333,7.5),dpi=dpi); fig.patch.set_facecolor(BG)
    bg=fig.add_axes([0,0,1,1]); bg.set_axis_off(); bg.set_xlim(0,1); bg.set_ylim(0,1)
    bg.add_patch(Rectangle((0,0.92),1,0.08,color=HDR,zorder=1))
    bg.text(0.015,0.965,"VENTAS QUANTITY DASHBOARD",color="white",fontsize=21,fontweight="bold",ha="left",va="center",zorder=3)
    bg.text(0.015,0.933,"%s    |    Report Date: %s    |    Last 7 Days    |    LOG SCALE"
            %(customer, G._short_date_long(report_date)),
            color="#B0C4DE",fontsize=10,ha="left",va="center",zorder=3)
    bg.text(0.985,0.96,"6D Technologies",color="white",fontsize=13,fontweight="bold",ha="right",va="center",zorder=3)
    M=0.015; colw=(1-4*M)/3; bottom=0.10; h=0.66
    for i,key in enumerate(ORDER):
        title,col,color=SPEC[key]
        vals=[data.get(key,{}).get(d,0.0) for d in dates]
        ax=fig.add_axes([M+i*(colw+M)+0.012, bottom, colw-0.024, h]); ax.set_facecolor(CHART_BG2)
        pos=[v for v in vals if v>0]
        for xi,v in zip(x,vals):
            if v>0:
                ax.bar(xi,v,0.6,color=color,zorder=3,alpha=0.95)
                ax.annotate(_compact(v),(xi,v),textcoords="offset points",xytext=(0,4),
                            ha="center",color="white",fontsize=8.5,fontweight="bold")
        ax.set_yscale("log"); ax.set_ylim(1, (max(pos)*3 if pos else 10))
        if not pos:
            ax.text(0.5,0.5,"No quantity data",transform=ax.transAxes,ha="center",va="center",
                    color="#9FB6D6",fontsize=11,style="italic")
        ax.set_xticks(x); ax.set_xticklabels([short(d) for d in dates],color="white",fontsize=8.5)
        ax.tick_params(axis="y",colors="#9FB6D6",labelsize=8); ax.tick_params(axis="x",colors="white")
        for s in ax.spines.values(): s.set_color("#24456f")
        ax.grid(axis="y",which="both",color="white",alpha=0.10,zorder=0)
        ax.set_title(title,color="white",fontsize=12,fontweight="bold",pad=8)
        ax.set_ylabel("Quantity (log)",color="#9FB6D6",fontsize=8)
    fig.savefig(str(out_png),dpi=dpi,facecolor=fig.get_facecolor()); plt.close(fig)
    print("Quantity PNG rendered:",out_png)
    return out_png

def _add_slide(prs, png):
    """Append the quantity slide (full-bleed PNG) to an existing Presentation."""
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(0x06,0x22,0x4A)
    slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
    return slide

def build(csv_paths, report_date, output, png=None):
    """Build a standalone quantity deck (one slide)."""
    data=_collect(csv_paths)
    tmp = Path(png) if png else (Path(output).with_suffix(".png"))
    render_png(data, report_date, tmp)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    _add_slide(prs, tmp)
    prs.save(output)
    print("Saved quantity deck:",output)
    return output

def append_to_pptx(pptx_path, csv_paths, report_date, png=None):
    """Render the quantity chart and append it as a new slide to an EXISTING .pptx
    (the main dashboard deck), producing a single merged file in place."""
    data=_collect(csv_paths)
    tmp = Path(png) if png else (Path(pptx_path).with_name("_qty_tmp.png"))
    render_png(data, report_date, tmp)
    prs=Presentation(str(pptx_path))
    _add_slide(prs, tmp)
    prs.save(str(pptx_path))
    print("Appended quantity slide to:", pptx_path)
    return pptx_path

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("--csv-file",nargs="+",required=True)
    ap.add_argument("--date",required=True)
    ap.add_argument("--output",default="",help="write a standalone one-slide quantity deck here")
    ap.add_argument("--append-to",default="",help="append the quantity slide to this existing .pptx")
    ap.add_argument("--png",default="")
    a=ap.parse_args()
    if a.append_to:
        append_to_pptx(a.append_to, a.csv_file, a.date, a.png or None)
    elif a.output:
        build(a.csv_file, a.date, a.output, a.png or None)
    else:
        ap.error("provide --output (standalone deck) or --append-to (merge into existing deck)")
